"""A claim batch as an Excel workbook in the payer layout (FEATURES 11.3).

The sheet is what the center sends to the payer: the center and payer header (contract,
claim number, period), one row per claimed service with the patient's card number, the
invoice, the service, the pre-approval reference and the claimed amount, then the total.
Once the payer answered, the accepted and rejected amounts and the payer's reasons follow.
Amounts are written as numbers with two decimals (``Decimal`` from the documents, never a
float computed here). Arabic sheets are right-to-left.
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from apps.claims import queries

__all__ = ["claim_workbook"]

_LABELS: dict[str, dict[str, str]] = {
    "en": {
        "title": "Insurance claim",
        "payer": "Payer",
        "contract": "Contract no.",
        "claim": "Claim no.",
        "period": "Period",
        "status": "Status",
        "no": "#",
        "patient": "Patient",
        "file_no": "File no.",
        "card": "Card no.",
        "invoice": "Invoice no.",
        "date": "Date",
        "code": "Service code",
        "service": "Service",
        "qty": "Qty",
        "gross": "Gross",
        "pre_approval": "Pre-approval",
        "claimed": "Claimed",
        "accepted": "Accepted",
        "rejected": "Rejected",
        "reason": "Payer reason",
        "total": "Total",
        "to": "to",
    },
    "ar": {
        "title": "مطالبة تأمين",
        "payer": "جهة التغطية",
        "contract": "رقم العقد",
        "claim": "رقم المطالبة",
        "period": "الفترة",
        "status": "الحالة",
        "no": "#",
        "patient": "المريض",
        "file_no": "رقم الملف",
        "card": "رقم البطاقة",
        "invoice": "رقم الفاتورة",
        "date": "التاريخ",
        "code": "رمز الخدمة",
        "service": "الخدمة",
        "qty": "الكمية",
        "gross": "الإجمالي",
        "pre_approval": "الموافقة المسبقة",
        "claimed": "المطالب به",
        "accepted": "المقبول",
        "rejected": "المرفوض",
        "reason": "سبب الجهة",
        "total": "الإجمالي",
        "to": "إلى",
    },
}

_STATUS: dict[str, dict[str, str]] = {
    "en": {
        "draft": "Draft",
        "submitted": "Submitted",
        "responded": "Answered",
        "closed": "Closed",
        "void": "Void",
    },
    "ar": {
        "draft": "مسودة",
        "submitted": "مُرسلة",
        "responded": "تم الرد",
        "closed": "مغلقة",
        "void": "ملغاة",
    },
}

_MONEY_FORMAT = "#,##0.00"


def _pick(row: dict[str, Any], field: str, language: str) -> str:
    ar, en = str(row[f"{field}_ar"]), str(row[f"{field}_en"])
    return (ar or en) if language == "ar" else (en or ar)


def claim_workbook(claim_id: int, language: str = "en") -> tuple[str, bytes]:
    """The claim batch ``claim_id`` as an ``.xlsx`` in ``language`` (ar or en), with its
    download name (the claim number, safe characters only)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    lang = "ar" if language == "ar" else "en"
    t = _LABELS[lang]
    data = queries.claim_print(claim_id)
    claim, payer, center = data["claim"], data["payer"], data["center"]
    answered = claim["status"] in ("responded", "closed")

    book = Workbook()
    sheet = book.active or book.create_sheet()
    sheet.title = claim["number"][:31]
    sheet.sheet_view.rightToLeft = lang == "ar"
    bold = Font(bold=True)

    sheet.append([_pick(center, "name", lang)])
    sheet.append([t["title"]])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"].font = Font(bold=True, size=12)
    header = [
        (t["payer"], _pick(payer, "name", lang)),
        (t["contract"], payer["contract_no"]),
        (t["claim"], claim["number"]),
        (
            t["period"],
            f"{claim['period_start'].isoformat()} {t['to']} {claim['period_end'].isoformat()}",
        ),
        (t["status"], _STATUS[lang][claim["status"]]),
    ]
    for label, value in header:
        sheet.append([label, value])
        sheet.cell(row=sheet.max_row, column=1).font = bold
    sheet.append([])

    columns = [
        t["no"],
        t["patient"],
        t["file_no"],
        t["card"],
        t["invoice"],
        t["date"],
        t["code"],
        t["service"],
        t["qty"],
        t["gross"],
        t["pre_approval"],
        t["claimed"],
    ]
    if answered:
        columns += [t["accepted"], t["rejected"], t["reason"]]
    sheet.append(columns)
    head_row = sheet.max_row
    for col in range(1, len(columns) + 1):
        cell = sheet.cell(row=head_row, column=col)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    money_cols = {10, 12} | ({13, 14} if answered else set())
    for index, line in enumerate(claim["lines"], start=1):
        row: list[Any] = [
            index,
            _pick(line["patient"], "full_name", lang),
            line["patient"]["file_no"],
            line["card_number"],
            line["invoice_number"],
            line["approved_on"],
            line["service_code"],
            _pick(line, "description", lang),
            line["quantity"],
            Decimal(line["gross"]),
            line["pre_approval_ref"],
            Decimal(line["amount_claimed"]),
        ]
        if answered:
            row += [
                Decimal(line["accepted_amount"]),
                Decimal(line["rejected_amount"]),
                line["payer_reason"],
            ]
        sheet.append(row)
        current = sheet.max_row
        sheet.cell(row=current, column=6).number_format = "yyyy-mm-dd"
        for col in money_cols:
            sheet.cell(row=current, column=col).number_format = _MONEY_FORMAT

    total: list[Any] = [None] * 11 + [Decimal(claim["claimed_total"])]
    total[0] = t["total"]
    if answered:
        total += [Decimal(claim["accepted_total"]), Decimal(claim["rejected_total"])]
    sheet.append(total)
    current = sheet.max_row
    for col in range(1, len(total) + 1):
        sheet.cell(row=current, column=col).font = bold
    for col in money_cols & set(range(1, len(total) + 1)):
        sheet.cell(row=current, column=col).number_format = _MONEY_FORMAT

    widths = [5, 28, 14, 16, 18, 12, 14, 30, 6, 14, 16, 14, 14, 14, 30]
    for col, width in enumerate(widths[: len(columns)], start=1):
        sheet.column_dimensions[sheet.cell(row=head_row, column=col).column_letter].width = width
    sheet.freeze_panes = sheet.cell(row=head_row + 1, column=1)

    out = io.BytesIO()
    book.save(out)
    name = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in claim["number"])
    return f"{name}-{lang}.xlsx", out.getvalue()
