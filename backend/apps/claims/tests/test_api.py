"""``/api/claims`` end to end through the HTTP API (FEATURES 11.2-11.7, invariant 7).

The accountant takes one payer's shares through the whole cycle: accrued, claimed, exported,
answered (accepted, partial, rejected), the rejection rebilled to the patient or written off,
a payer payment allocated per claim, aging reduced. After every step the payer receivable
derived from the documents equals the AR_PAYER ledger balance, and no step before the payer
payment touches cash or bank (invariant 7).
"""

from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

import pytest
from django.utils import timezone

from apps.claims import services as cs
from apps.claims.models import ClaimLine, PayerPayment
from apps.ledger import services as ledger
from apps.payments import services as payments
from apps.payments.tests import api_kit as kit
from apps.payments.tests import fin
from conftest import TEST_PASSWORD

pytestmark = pytest.mark.django_db

D = Decimal
ok, error = kit.ok, kit.error


def _money() -> tuple[Decimal, Decimal, Decimal]:
    return (
        ledger.account_balance("CASH"),
        ledger.account_balance("BANK"),
        ledger.account_balance("BANK_PENDING"),
    )


def _reconciles(payer: Any) -> None:
    row = cs.payer_receivables(payer=payer).get(payer.pk)
    documents = row.receivable if row else D("0")
    assert documents == ledger.account_balance("AR_PAYER", payer=payer)


class Setup:
    def __init__(self, make_user: Any) -> None:
        self.desk = kit.desk(make_user)
        self.acc = self.desk.accountant
        self.payer = fin.payer(percent="70")
        self.patient = fin.patient(full_name_en="Amna Yousif")
        self.visit = fin.visit(self.patient, payer_obj=self.payer, card_number="AM-77")
        doctor = self.desk.doctor.user
        lines = fin.order(
            self.visit,
            doctor,
            fin.priced("lab", "10000.00"),
            fin.priced("lab", "5000.00"),
            fin.priced("procedure", "2000.00"),
        )
        self.invoice = fin.invoice(self.visit, self.desk.cashier.user, lines)
        self.il = list(self.invoice.lines.order_by("line_no"))


@pytest.fixture
def s(make_user: Any) -> Setup:
    return Setup(make_user)


def _url(path: str) -> str:
    return f"/api/claims{path}"


#: The manager of ``kit.desk`` approves rebills and write-offs as the second person (ADR 0018).
SECOND = {"approver": {"username": "mgr", "password": TEST_PASSWORD}}


def test_receivables_accrue_and_a_claim_is_built_from_chosen_lines(s: Setup) -> None:
    acc = s.acc.api
    before = _money()
    options = ok(acc.get(_url("/options")))
    assert s.payer.pk in {p["id"] for p in options["payers"]}
    assert {r["code"] for r in options["write_off_reasons"]} >= {"NOT_COVERED", "OTHER"}
    assert options["open_shift"] is None

    rec = ok(acc.get(_url("/receivables")))
    row = next(r for r in rec["items"] if r["payer"]["id"] == s.payer.pk)
    assert row["stages"]["accrued"] == "11900.00"  # 7,000 + 3,500 + 1,400
    assert row["stages"]["receivable"] == "11900.00"
    assert row["stages"]["collected"] == "0.00"
    assert row["aging"]["days_0_30"] == "11900.00"

    today = timezone.localdate().isoformat()
    accrued = ok(
        acc.get(_url(f"/accrued?payer_id={s.payer.pk}&period_start={today}&period_end={today}"))
    )
    assert [i["amount"] for i in accrued["items"]] == ["7000.00", "3500.00", "1400.00"]
    first = accrued["items"][0]
    assert first["card_number"] == "AM-77"
    assert first["invoice_number"] == s.invoice.number
    assert first["patient"]["full_name_en"] == "Amna Yousif"
    assert accrued["total"] == "11900.00"

    body = {
        "payer_id": s.payer.pk,
        "period_start": today,
        "period_end": today,
        "invoice_line_ids": [s.il[0].pk, s.il[1].pk],
        "note": "October",
    }
    claim = ok(acc.post(_url("/batches"), body), 201)
    assert claim["status"] == "draft"
    assert claim["claimed_total"] == "10500.00"
    assert [ln["stage"] for ln in claim["lines"]] == ["claimed", "claimed"]
    # The third share stays accrued, claimable in a later batch.
    accrued = ok(acc.get(_url(f"/accrued?payer_id={s.payer.pk}")))
    assert [i["invoice_line_id"] for i in accrued["items"]] == [s.il[2].pk]
    error(acc.post(_url("/batches"), body), 409, "CLAIM_LINE_NOT_ACCRUED")

    removed = ok(
        acc.request("DELETE", _url(f"/batches/{claim['id']}/lines/{claim['lines'][1]['id']}"))
    )
    assert removed["claimed_total"] == "7000.00"
    assert len(removed["lines"]) == 1
    listed = ok(acc.get(_url(f"/batches?payer_id={s.payer.pk}&status=draft")))
    assert [c["number"] for c in listed["items"]] == [claim["number"]]
    # Claiming is not collecting.
    assert _money() == before
    _reconciles(s.payer)


def _submitted(s: Setup) -> dict[str, Any]:
    today = timezone.localdate().isoformat()
    claim = ok(
        s.acc.api.post(
            _url("/batches"), {"payer_id": s.payer.pk, "period_start": today, "period_end": today}
        ),
        201,
    )
    return ok(s.acc.api.post(_url(f"/batches/{claim['id']}/submit")))


def test_export_writes_typed_text_never_as_a_formula(s: Setup) -> None:
    from openpyxl import load_workbook

    s.patient.full_name_en = '=HYPERLINK("http://x","y")'
    s.patient.save(update_fields=["full_name_en"])
    claim = _submitted(s)
    answers = {
        "responses": [
            {"claim_line_id": ln["id"], "outcome": "rejected", "reason": "=1+1"}
            for ln in claim["lines"]
        ]
    }
    ok(s.acc.api.post(_url(f"/batches/{claim['id']}/responses"), answers))
    response = s.acc.api.get(_url(f"/batches/{claim['id']}/export?language=en"))
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    cells = [c for row in sheet.iter_rows() for c in row if c.value is not None]
    typed = [c for c in cells if str(c.value).startswith("=")]
    assert {str(c.value) for c in typed} == {'=HYPERLINK("http://x","y")', "=1+1"}
    assert all(c.data_type == "s" for c in typed)


def test_export_is_a_valid_workbook_in_the_payer_layout(s: Setup) -> None:
    from openpyxl import load_workbook

    claim = _submitted(s)
    for language in ("en", "ar"):
        response = s.acc.api.get(_url(f"/batches/{claim['id']}/export?language={language}"))
        assert response.status_code == 200, response.content
        assert response["Content-Type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert claim["number"] in response["Content-Disposition"]
        book = load_workbook(io.BytesIO(response.content))
        sheet = book.active
        assert sheet is not None
        values = [list(row) for row in sheet.iter_rows(values_only=True)]
        flat = [str(v) for row in values for v in row if v is not None]
        assert claim["number"] in flat
        assert s.payer.code in flat or s.payer.name_en in flat or s.payer.name_ar in flat
        assert "AM-77" in flat
        assert s.invoice.number in flat
        assert sheet.sheet_view.rightToLeft is (language == "ar")
        # Amounts are numbers, and the total row adds up to the claimed total.
        claimed = [row[11] for row in values if isinstance(row[0], int)]
        assert claimed == [D("7000.00"), D("3500.00"), D("1400.00")]
        assert values[-1][11] == D("11900.00")

    printed = ok(s.acc.api.get(_url(f"/batches/{claim['id']}/print")))
    assert printed["payer"]["id"] == s.payer.pk
    assert printed["claim"]["number"] == claim["number"]
    assert printed["center"]["name_en"] is not None


def test_answers_rebill_write_off_and_payment_by_claim(s: Setup) -> None:
    acc = s.acc.api
    claim = _submitted(s)
    cid = claim["id"]
    full, part, rejected = (ln["id"] for ln in claim["lines"])
    before = _money()

    answers = {
        "responses": [
            {"claim_line_id": full, "outcome": "accepted", "reference": "RA-1"},
            {"claim_line_id": part, "outcome": "partial", "accepted": "2000"},
            {"claim_line_id": rejected, "outcome": "rejected", "reason": "not covered"},
        ]
    }
    error(acc.post(_url(f"/batches/{cid}/responses"), answers), 409, "REASON_REQUIRED")
    answers["responses"][1]["reason"] = "tariff limit"
    answers["responses"][1]["accepted"] = "3500"
    error(acc.post(_url(f"/batches/{cid}/responses"), answers), 409, "CLAIM_PARTIAL_INVALID")
    answers["responses"][1]["accepted"] = "2000"
    answered = ok(acc.post(_url(f"/batches/{cid}/responses"), answers))
    assert answered["status"] == "responded"
    assert [ln["stage"] for ln in answered["lines"]] == [
        "accepted",
        "partially_accepted",
        "rejected",
    ]
    assert (answered["accepted_total"], answered["rejected_total"]) == ("9000.00", "2900.00")
    assert answered["lines"][1]["unresolved_rejection"] == "1500.00"
    error(acc.post(_url(f"/batches/{cid}/close")), 409, "CLAIM_NOT_SETTLED")

    # Rebill the rejected part: the patient owes it on the original invoice.
    invoice_url = f"/api/billing/invoices/{s.invoice.pk}"
    owed_before = D(ok(acc.get(invoice_url))["outstanding"])
    error(
        acc.post(
            _url(f"/batches/{cid}/lines/{full}/resolve"),
            {**SECOND, "resolution": "rebilled", "reason": "NOT_COVERED"},
        ),
        409,
        "CLAIM_NOTHING_REJECTED",
    )
    rebilled = ok(
        acc.post(
            _url(f"/batches/{cid}/lines/{rejected}/resolve"),
            {
                **SECOND,
                "resolution": "rebilled",
                "reason": "NOT_COVERED",
                "note": "patient informed",
            },
        )
    )
    line = rebilled["lines"][2]
    assert (line["stage"], line["resolution"]) == ("rebilled", "rebilled")
    assert line["resolution_reason"]["code"] == "NOT_COVERED"
    assert line["resolved_by"]["username"] == s.acc.username
    assert D(ok(acc.get(invoice_url))["outstanding"]) == owed_before + D("1400.00")
    # Write off the partial rejection; an "other" reason needs a note.
    error(
        acc.post(
            _url(f"/batches/{cid}/lines/{part}/resolve"),
            {**SECOND, "resolution": "written_off", "reason": "OTHER"},
        ),
        409,
        "REASON_NOTE_REQUIRED",
    )
    ok(
        acc.post(
            _url(f"/batches/{cid}/lines/{part}/resolve"),
            {**SECOND, "resolution": "written_off", "reason": "SMALL_BALANCE"},
        )
    )
    # Nothing so far moved cash or bank money.
    assert _money() == before
    _reconciles(s.payer)

    aging = ok(acc.get(_url("/aging")))
    row = next(r for r in aging["items"] if r["payer"]["id"] == s.payer.pk)
    assert row["aging"]["total"] == "9000.00"

    payable = ok(acc.get(_url(f"/payers/{s.payer.pk}/payable")))
    assert [(c["id"], c["unpaid_total"]) for c in payable] == [(cid, "9000.00")]
    options = ok(acc.get(_url("/options")))
    bank = options["banks"][0]["id"]
    payment_body = {
        "payer_id": s.payer.pk,
        "amount": "8000",
        "method": "bank_transfer",
        "bank_id": bank,
        "reference": "RA-2026-10",
        "received_on": timezone.localdate().isoformat(),
        "claims": [{"claim_id": cid, "amount": "8000"}],
    }
    error(
        acc.post(_url("/payer-payments"), {**payment_body, "amount": "9000"}),
        409,
        "PAYER_PAYMENT_UNBALANCED",
    )
    error(
        acc.post(_url("/payer-payments"), {**payment_body, "method": "cash"}),
        409,
        "SHIFT_NOT_OPEN",
    )
    paid = ok(acc.post(_url("/payer-payments"), payment_body), 201)
    assert paid["standing"] == "bank"
    assert [(a["claim_number"], a["amount"]) for a in paid["allocations"]] == [
        (claim["number"], "7000.00"),
        (claim["number"], "1000.00"),
    ]
    assert _money()[1] - before[1] == D("8000.00")
    aging = ok(acc.get(_url("/aging")))
    row = next(r for r in aging["items"] if r["payer"]["id"] == s.payer.pk)
    assert row["aging"]["total"] == "1000.00"
    _reconciles(s.payer)

    # The payer short-pays the last 1,000: written off with a reason, then the claim closes.
    detail = ok(acc.get(_url(f"/batches/{cid}")))
    assert detail["paid_total"] == "8000.00"
    assert [p["number"] for p in detail["payments"]] == [paid["number"]]
    ok(
        acc.post(
            _url(f"/batches/{cid}/lines/{part}/write-off"),
            {**SECOND, "amount": "1000", "reason": "SMALL_BALANCE"},
        )
    )
    closed = ok(acc.post(_url(f"/batches/{cid}/close")))
    assert closed["status"] == "closed"
    assert closed["receivable"] == "0.00"
    assert ledger.account_balance("AR_PAYER", payer=s.payer) == D("0.00")
    listed = ok(acc.get(_url(f"/payer-payments?payer_id={s.payer.pk}")))
    assert [p["number"] for p in listed["items"]] == [paid["number"]]


def test_cheques_cash_and_reversal(s: Setup, make_user: Any) -> None:
    acc = s.acc.api
    claim = _submitted(s)
    cid = claim["id"]
    ok(
        acc.post(
            _url(f"/batches/{cid}/responses"),
            {
                "responses": [
                    {"claim_line_id": ln["id"], "outcome": "accepted"} for ln in claim["lines"]
                ]
            },
        )
    )
    today = timezone.localdate().isoformat()
    cheque = ok(
        acc.post(
            _url("/payer-payments"),
            {
                "payer_id": s.payer.pk,
                "amount": "7000",
                "method": "cheque",
                "received_on": today,
            },
        ),
        201,
    )
    assert cheque["standing"] == "cheque_pending"
    pending = ok(acc.get(_url("/payer-payments?standing=cheque_pending")))
    assert [p["id"] for p in pending["items"]] == [cheque["id"]]
    assert ledger.account_balance("BANK_PENDING") == D("7000.00")
    cleared = ok(acc.post(_url(f"/payer-payments/{cheque['id']}/clear"), {"note": "statement"}))
    assert cleared["standing"] == "cheque_cleared"
    error(acc.post(_url(f"/payer-payments/{cheque['id']}/clear"), {}), 409, "PAYMENT_NOT_PENDING")
    error(acc.post(_url(f"/payer-payments/{cheque['id']}/reverse"), {}), 409, "REASON_REQUIRED")
    reversed_ = ok(acc.post(_url(f"/payer-payments/{cheque['id']}/reverse"), {"note": "bounced"}))
    assert reversed_["standing"] == "reversed"
    assert ok(acc.get(_url(f"/batches/{cid}")))["paid_total"] == "0.00"
    _reconciles(s.payer)

    # Payer cash goes into the recorder's open shift and counts in its drawer.
    admin = kit.actor(make_user, "boss", "admin")
    shift = payments.open_shift(admin.user, D("0.00"))
    options = ok(admin.api.get(_url("/options")))
    assert options["open_shift"]["number"] == shift.number
    cash = ok(
        admin.api.post(
            _url("/payer-payments"),
            {"payer_id": s.payer.pk, "amount": "1400", "method": "cash", "received_on": today},
        ),
        201,
    )
    assert cash["standing"] == "cash"
    assert cash["shift"]["id"] == shift.pk
    assert payments.expected_cash(shift) == D("1400.00")
    error(
        admin.api.post(_url(f"/payer-payments/{cash['id']}/reverse"), {"note": "x"}),
        409,
        "PAYMENT_NOT_REVERSIBLE",
    )
    assert PayerPayment.objects.count() == 2
    _reconciles(s.payer)


def test_paths_and_references_are_checked(s: Setup) -> None:
    acc = s.acc.api
    claim = _submitted(s)
    other = fin.payer(percent="50")
    error(acc.get(_url("/batches/999999")), 404, "NOT_FOUND")
    # A line of another claim is never reached through this claim's path.
    line = ClaimLine.objects.get(pk=claim["lines"][0]["id"])
    assert line.claim_id == claim["id"]
    error(
        acc.post(
            _url(f"/batches/{claim['id'] + 1000}/lines/{line.pk}/resolve"),
            {**SECOND, "resolution": "rebilled", "reason": "NOT_COVERED"},
        ),
        404,
        "NOT_FOUND",
    )
    error(
        acc.post(
            _url(f"/batches/{claim['id']}/responses"),
            {"responses": [{"claim_line_id": 999_999, "outcome": "accepted"}]},
        ),
        409,
        "CLAIM_LINE_UNKNOWN",
    )
    error(
        acc.post(
            _url("/payer-payments"),
            {
                "payer_id": other.pk,
                "amount": "10",
                "method": "cheque",
                "received_on": timezone.localdate().isoformat(),
                "claims": [{"claim_id": claim["id"], "amount": "10"}],
            },
        ),
        409,
        "CLAIM_NOTHING_UNPAID",
    )
    error(
        acc.post(
            _url("/payer-payments"),
            {
                "payer_id": s.payer.pk,
                "amount": "10",
                "method": "cheque",
                "received_on": timezone.localdate().isoformat(),
                "lines": [
                    {"claim_line_id": claim["lines"][0]["id"], "amount": "5"},
                    {"claim_line_id": claim["lines"][0]["id"], "amount": "5"},
                ],
            },
        ),
        409,
        "DUPLICATE_LINE",
    )
    voided = ok(acc.post(_url(f"/batches/{claim['id']}/void"), {"note": "sent twice"}))
    assert voided["status"] == "void"
    assert ok(acc.get(_url(f"/accrued?payer_id={s.payer.pk}")))["total"] == "11900.00"


def test_manager_reads_but_does_not_act(s: Setup) -> None:
    claim = _submitted(s)
    manager = s.desk.manager.api
    ok(manager.get(_url("/receivables")))
    ok(manager.get(_url(f"/batches/{claim['id']}")))
    body = {"responses": [{"claim_line_id": claim["lines"][0]["id"], "outcome": "accepted"}]}
    response = manager.post(_url(f"/batches/{claim['id']}/responses"), body)
    assert response.status_code == 403
    assert response.json()["details"]["permission"] == "claims.record_response"


def test_rebill_needs_a_second_persons_credentials(s: Setup) -> None:
    """ADR 0018: the options say a second person approves; the endpoint refuses the
    recorder alone and the recorder's own credentials, and records the approver."""
    acc = s.acc.api
    assert ok(acc.get(_url("/options")))["second_approver_required"] is True
    claim = _submitted(s)
    line_id = claim["lines"][0]["id"]
    answers = {"responses": [{"claim_line_id": line_id, "outcome": "rejected", "reason": "x"}]}
    ok(acc.post(_url(f"/batches/{claim['id']}/responses"), answers))
    path = _url(f"/batches/{claim['id']}/lines/{line_id}/resolve")
    body: dict[str, Any] = {"resolution": "written_off", "reason": "NOT_COVERED"}
    error(acc.post(path, body), 409, "SECOND_APPROVER_REQUIRED")
    own = {"username": s.acc.user.username, "password": TEST_PASSWORD}
    error(acc.post(path, {**body, "approver": own}), 409, "SECOND_APPROVER_REQUIRED")
    wrong = {"username": "mgr", "password": "not-it"}
    error(acc.post(path, {**body, "approver": wrong}), 409, "APPROVER_INVALID")
    cashier = {"username": s.desk.cashier.user.username, "password": TEST_PASSWORD}
    error(acc.post(path, {**body, "approver": cashier}), 409, "APPROVER_NOT_PERMITTED")
    done = ok(acc.post(path, {**body, **SECOND}))
    assert done["lines"][0]["resolution"] == "written_off"
    assert ClaimLine.objects.get(pk=line_id).resolution_approved_by == s.desk.manager.user
