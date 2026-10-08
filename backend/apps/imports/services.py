"""Excel import of patients with preview, validation and duplicate detection (FEATURES 1.8).

1. :func:`preview_patients` reads the uploaded sheet (``.xlsx``, or ``.csv`` saved from
   Excel), matches its header row, reads every row with ``domain.patient_import`` and looks
   for duplicates: against existing files with ``patients.services.find_duplicates`` and
   against earlier rows of the same file (same phone or national ID). It stores an
   ``ImportJob`` (``validated``) with one ``ImportRow`` per sheet row. Nothing is registered.
2. :func:`confirm_job` registers the valid rows through ``patients.services.register_patient``
   (which checks duplicates again, so a file registered since the preview is not created
   twice) and, only when asked, the rows flagged as possible duplicates. Rows with errors are
   never imported. Each row records the file it created or why it was skipped.
3. :func:`cancel_job` drops a preview.

Row ``errors`` are ``{"code", "field"}`` and ``warnings`` are duplicate hints
``{"code": "phone"|"national_id"|"name_dob", "patient_id", "file_no", "name"}`` or
``{"code": "in_file", "row_no"}``.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterator
from datetime import date
from pathlib import Path
from typing import Any

import pghistory
from django.core.files.base import ContentFile
from django.db import connection, transaction
from django.db.models import QuerySet
from django.utils import timezone

from apps.core.models import User
from apps.core.services import require_permission
from apps.imports.models import ImportJob, ImportKind, ImportRow, ImportStatus, RowStatus
from apps.patients import services as patient_services
from domain import patient_import as dpi
from domain.errors import DomainError

__all__ = [
    "MAX_FILE_BYTES",
    "MAX_ROWS",
    "cancel_job",
    "confirm_job",
    "job_rows",
    "patient_template",
    "preview_patients",
]

#: Largest accepted upload (a 2,000-row sheet is far below this).
MAX_FILE_BYTES = 5 * 1024 * 1024
#: Rows per import; larger lists are split into several files.
MAX_ROWS = 2000

_XLSX = (".xlsx", ".xlsm")
_CSV = (".csv",)


# --- reading the file -----------------------------------------------------------------------


def _xlsx_rows(content: bytes) -> Iterator[tuple[Any, ...]]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise DomainError("IMPORT_FILE_INVALID", "The file is not a readable Excel file") from exc
    try:
        sheet = book.worksheets[0]
        yield from sheet.iter_rows(values_only=True)
    finally:
        book.close()


def _csv_rows(content: bytes) -> Iterator[list[str]]:
    for encoding in ("utf-8-sig", "cp1256"):
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        yield from csv.reader(io.StringIO(text))
        return
    raise DomainError("IMPORT_FILE_INVALID", "The file is not readable text")  # pragma: no cover


def _sheet_rows(filename: str, content: bytes) -> Iterator[tuple[int, list[Any]]]:
    """``(sheet row number, cells)`` of the first sheet, 1-based like Excel."""
    suffix = Path(filename).suffix.lower()
    if suffix in _XLSX:
        rows: Iterator[Any] = _xlsx_rows(content)
    elif suffix in _CSV:
        rows = _csv_rows(content)
    else:
        raise DomainError(
            "IMPORT_FILE_INVALID", "Upload an Excel (.xlsx) or CSV file", filename=filename
        )
    for number, cells in enumerate(rows, start=1):
        yield number, list(cells)


def _blank(cells: list[Any]) -> bool:
    return all(c is None or str(c).strip() == "" for c in cells)


# --- normalization in bulk ------------------------------------------------------------------


def _normalized_phones(values: list[str]) -> list[str]:
    """``hs_normalize_phone`` of every value in one query (the stored ``phone_norm`` folding)."""
    if not values:
        return []
    with connection.cursor() as cursor:
        cursor.execute(
            "SELECT hs_normalize_phone(v) FROM unnest(%s::text[]) WITH ORDINALITY AS t(v, n) "
            "ORDER BY n",
            [values],
        )
        return [str(r[0] or "") for r in cursor.fetchall()]


# --- preview --------------------------------------------------------------------------------


def _json_data(data: dict[str, Any]) -> dict[str, Any]:
    return {k: (v.isoformat() if isinstance(v, date) else v) for k, v in data.items()}


def _patient_name(patient: Any) -> str:
    return str(patient.full_name_ar or patient.full_name_en)


def preview_patients(
    *, filename: str, content: bytes, actor: User, today: date | None = None
) -> ImportJob:
    """Validate a patient sheet and store the preview (FEATURES 1.8). Registers nothing.

    Raises:
        PermissionRequired: the actor lacks ``imports.run``.
        DomainError: ``IMPORT_FILE_TOO_LARGE``, ``IMPORT_FILE_INVALID``,
            ``IMPORT_HEADERS_MISSING`` (``details.missing``), ``IMPORT_NO_ROWS``,
            ``IMPORT_TOO_MANY_ROWS`` (``details.limit``).
    """
    require_permission(actor, "imports.run")
    if len(content) > MAX_FILE_BYTES:
        raise DomainError(
            "IMPORT_FILE_TOO_LARGE", "The file is too large", limit_mb=MAX_FILE_BYTES // 2**20
        )
    on = today or timezone.localdate()
    rows = _sheet_rows(filename, content)
    header: dict[int, str] | None = None
    parsed: list[tuple[int, dpi.ParsedRow]] = []
    for number, cells in rows:
        if header is None:
            if _blank(cells):
                continue
            header = dpi.match_headers(cells)
            missing = dpi.missing_headers(header)
            if missing:
                raise DomainError(
                    "IMPORT_HEADERS_MISSING", "Required columns are missing", missing=missing
                )
            continue
        row = dpi.parse_row(
            {key: cells[pos] if pos < len(cells) else None for pos, key in header.items()},
            today=on,
        )
        if row is None:
            continue
        parsed.append((number, row))
        if len(parsed) > MAX_ROWS:
            raise DomainError("IMPORT_TOO_MANY_ROWS", "Too many rows in one file", limit=MAX_ROWS)
    if header is None or not parsed:
        raise DomainError("IMPORT_NO_ROWS", "The file has no patient rows")

    # Repeats inside the file: the same phone (either number) or national ID.
    raw_phones = [r.data[k] for _, r in parsed for k in ("phone", "phone_alt")]
    folded = iter(_normalized_phones(raw_phones))
    keys: list[tuple[int, set[str]]] = []
    for number, row in parsed:
        row_keys = {f"phone:{p}" for p in (next(folded), next(folded)) if len(p) >= 6}
        if row.data["national_id"]:
            row_keys.add(f"nid:{row.data['national_id']}")
        keys.append((number, row_keys))
    repeats = dpi.repeated_rows(keys)

    import_rows: list[ImportRow] = []
    counts = {RowStatus.VALID: 0, RowStatus.ERROR: 0, RowStatus.DUPLICATE: 0}
    for number, row in parsed:
        errors = [
            {"code": c, "field": f} for c, f in zip(row.errors, row.error_fields, strict=True)
        ]
        warnings: list[dict[str, Any]] = []
        duplicate_of: int | None = None
        if not errors:
            candidates = patient_services.find_duplicates(
                full_name_ar=row.data["full_name_ar"],
                full_name_en=row.data["full_name_en"],
                phone=row.data["phone"],
                phone_alt=row.data["phone_alt"],
                date_of_birth=row.data["date_of_birth"],
                national_id=row.data["national_id"],
                limit=3,
            )
            for cand in candidates:
                for reason in cand.reasons:
                    warnings.append(
                        {
                            "code": reason,
                            "patient_id": cand.patient.pk,
                            "file_no": cand.patient.file_no,
                            "name": _patient_name(cand.patient),
                        }
                    )
            if candidates:
                duplicate_of = candidates[0].patient.pk
            if number in repeats:
                warnings.append({"code": "in_file", "row_no": repeats[number]})
        status = RowStatus.ERROR if errors else RowStatus.DUPLICATE if warnings else RowStatus.VALID
        counts[status] += 1
        import_rows.append(
            ImportRow(
                row_no=number,
                status=status,
                data=_json_data(row.data),
                errors=errors,
                warnings=warnings,
                duplicate_of_id=duplicate_of,
            )
        )

    with transaction.atomic(), pghistory.context(user=actor.pk, reason="patient import preview"):
        job = ImportJob(
            kind=ImportKind.PATIENTS,
            status=ImportStatus.VALIDATED,
            original_filename=Path(filename).name[:255],
            total_rows=len(import_rows),
            valid_rows=counts[RowStatus.VALID],
            error_rows=counts[RowStatus.ERROR],
            duplicate_rows=counts[RowStatus.DUPLICATE],
            uploaded_by=actor,
            validated_at=timezone.now(),
        )
        job.file.save(Path(filename).name, ContentFile(content), save=False)
        job.save()
        for r in import_rows:
            r.job = job
        ImportRow.objects.bulk_create(import_rows)
    return job


# --- confirm and cancel ---------------------------------------------------------------------


def _locked_open_job(job: ImportJob) -> ImportJob:
    locked = ImportJob.objects.select_for_update().get(pk=job.pk)
    if locked.status != ImportStatus.VALIDATED:
        raise DomainError(
            "IMPORT_JOB_CLOSED",
            "This import was already confirmed or cancelled",
            status=locked.status,
        )
    return locked


def _patient_data(data: dict[str, Any]) -> patient_services.PatientData:
    dob = data.get("date_of_birth")
    return patient_services.PatientData(
        sex=str(data.get("sex") or ""),
        full_name_ar=str(data.get("full_name_ar") or ""),
        full_name_en=str(data.get("full_name_en") or ""),
        date_of_birth=date.fromisoformat(dob) if dob else None,
        age_years=data.get("age_years"),
        phone=str(data.get("phone") or ""),
        phone_alt=str(data.get("phone_alt") or ""),
        address=str(data.get("address") or ""),
        national_id=str(data.get("national_id") or ""),
        emergency_contact_name=str(data.get("emergency_contact_name") or ""),
        emergency_contact_phone=str(data.get("emergency_contact_phone") or ""),
        notes=str(data.get("notes") or ""),
    )


def confirm_job(job: ImportJob, *, actor: User, include_duplicates: bool = False) -> ImportJob:
    """Register the previewed patients (FEATURES 1.8).

    Valid rows are registered with the normal duplicate check, so a matching file created
    since the preview skips the row (``DUPLICATE_PATIENT``). Possible duplicates are
    registered only with ``include_duplicates`` (the supervisor looked at them). Rows with
    errors never are.

    Raises:
        PermissionRequired: the actor lacks ``imports.run``.
        DomainError: ``IMPORT_JOB_CLOSED``.
    """
    require_permission(actor, "imports.run")
    with transaction.atomic(), pghistory.context(user=actor.pk, reason="patient import"):
        locked = _locked_open_job(job)
        wanted = [RowStatus.VALID, *([RowStatus.DUPLICATE] if include_duplicates else [])]
        imported = skipped = 0
        for row in locked.rows.select_for_update().order_by("row_no"):
            if row.status not in wanted:
                if row.status != RowStatus.ERROR:
                    row.status = RowStatus.SKIPPED
                    row.save(update_fields=["status"])
                    skipped += 1
                continue
            try:
                patient = patient_services.register_patient(
                    _patient_data(row.data),
                    actor=actor,
                    confirm_not_duplicate=row.status == RowStatus.DUPLICATE,
                )
            except DomainError as exc:
                row.status = RowStatus.SKIPPED
                row.errors = [{"code": exc.code, "field": ""}]
                row.save(update_fields=["status", "errors"])
                skipped += 1
                continue
            row.status = RowStatus.IMPORTED
            row.result_id = patient.pk
            row.save(update_fields=["status", "result_id"])
            imported += 1
        locked.status = ImportStatus.CONFIRMED
        locked.imported_rows = imported
        locked.summary = {"skipped": skipped, "include_duplicates": include_duplicates}
        locked.confirmed_by = actor
        locked.confirmed_at = timezone.now()
        locked.save()
    return locked


def cancel_job(job: ImportJob, *, actor: User) -> ImportJob:
    """Drop a preview that will not be imported.

    Raises:
        DomainError: ``IMPORT_JOB_CLOSED``.
    """
    require_permission(actor, "imports.run")
    with transaction.atomic(), pghistory.context(user=actor.pk, reason="cancel import"):
        locked = _locked_open_job(job)
        locked.status = ImportStatus.CANCELLED
        locked.save(update_fields=["status"])
    return locked


def job_rows(job: ImportJob, *, status: str | None = None) -> QuerySet[ImportRow]:
    """The job's rows in sheet order; ``status="problems"`` lists errors and duplicates."""
    rows = job.rows.all()
    if status == "problems":
        rows = rows.filter(status__in=[RowStatus.ERROR, RowStatus.DUPLICATE])
    elif status:
        rows = rows.filter(status=status)
    return rows.order_by("row_no")


def patient_template(language: str = "en") -> bytes:
    """An empty ``.xlsx`` with the import header row in ``language`` (ar or en)."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active or book.create_sheet()
    sheet.title = "patients"
    labels = [c.label_ar if language == "ar" else c.label_en for c in dpi.COLUMNS]
    sheet.append(labels)
    if language == "ar":
        sheet.sheet_view.rightToLeft = True
    for index, label in enumerate(labels, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = max(
            14, len(label) + 4
        )
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
