"""Excel import of patients (FEATURES 1.8): preview, validation, duplicates, confirm, API."""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from api.errors import PermissionRequired
from apps.imports import services as imp
from apps.imports.models import ImportJob, ImportRow
from apps.patients import services as ps
from apps.patients.models import Patient
from conftest import ApiClient
from domain.errors import DomainError

pytestmark = pytest.mark.django_db

HEADER = ["Name (Arabic)", "Name (English)", "Sex", "Date of birth", "Age (years)", "Phone"]


@pytest.fixture(autouse=True)
def _media(settings: Any, tmp_path: Any) -> None:
    settings.MEDIA_ROOT = str(tmp_path)


@pytest.fixture
def admin(make_user):
    return make_user("importer", roles=["admin"])


@pytest.fixture
def clerk(make_user):
    return make_user("clerk", roles=["receptionist"])


def _xlsx(rows: list[list[Any]], header: list[str] = HEADER) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _rows(job: ImportJob) -> dict[int, ImportRow]:
    return {r.row_no: r for r in job.rows.all()}


def test_preview_validates_rows_and_flags_duplicates(admin, clerk) -> None:
    existing = ps.register_patient(
        ps.PatientData(sex="male", full_name_ar="عثمان موسى", phone="0912777001"),
        actor=clerk,
    )
    content = _xlsx(
        [
            ["أحمد علي", "", "ذكر", date(1990, 1, 2), None, "0912777100"],
            ["", "", "female", None, 30, "0912777101"],  # no name
            ["Sara", "Sara Ali", "unknown", None, None, ""],  # bad sex
            ["عثمان موسى", "", "M", None, 50, "+249912777001"],  # phone of an existing file
            ["علي أحمد", "", "male", "2099-01-01", None, ""],  # future birth date
            ["أمنة", "", "F", None, None, "0912777100"],  # same phone as row 2
            [None, None, None, None, None, None],  # blank: skipped
        ]
    )
    job = imp.preview_patients(filename="patients.xlsx", content=content, actor=admin)
    assert job.status == "validated"
    assert (job.total_rows, job.valid_rows, job.error_rows, job.duplicate_rows) == (6, 1, 3, 2)
    rows = _rows(job)
    assert rows[2].status == "valid"
    assert rows[2].data["date_of_birth"] == "1990-01-02"
    assert rows[3].errors == [{"code": "NAME_REQUIRED", "field": "full_name_ar"}]
    assert rows[4].errors == [{"code": "INVALID_SEX", "field": "sex"}]
    assert rows[5].status == "duplicate"
    assert rows[5].duplicate_of_id == existing.pk
    assert rows[5].warnings[0]["code"] == "phone"
    assert rows[5].warnings[0]["file_no"] == existing.file_no
    assert rows[6].errors[0]["code"] == "INVALID_DATE_OF_BIRTH"
    assert rows[7].status == "duplicate"
    assert rows[7].warnings == [{"code": "in_file", "row_no": 2}]
    assert 8 not in rows
    # Nothing is registered by a preview.
    assert Patient.objects.count() == 1


def test_confirm_registers_valid_rows_and_duplicates_only_when_asked(admin, clerk) -> None:
    ps.register_patient(
        ps.PatientData(sex="male", full_name_ar="عثمان موسى", phone="0912777001"), actor=clerk
    )
    content = _xlsx(
        [
            ["أحمد علي", "", "ذكر", None, 40, "0912777100"],
            ["عثمان موسى", "", "M", None, 50, "0912777001"],
            ["", "", "F", None, None, ""],
        ]
    )
    job = imp.preview_patients(filename="p.xlsx", content=content, actor=admin)
    done = imp.confirm_job(job, actor=admin)
    assert (done.status, done.imported_rows, done.confirmed_by) == ("confirmed", 1, admin)
    rows = _rows(done)
    assert rows[2].status == "imported"
    assert rows[2].result_id is not None
    created = Patient.objects.get(pk=rows[2].result_id)
    assert (created.full_name_ar, created.sex, created.dob_is_estimated) == (
        "أحمد علي",
        "male",
        True,
    )
    assert created.created_by == admin
    assert rows[3].status == "skipped"
    assert rows[4].status == "error"
    assert done.summary["skipped"] == 1
    with pytest.raises(DomainError) as exc:
        imp.confirm_job(done, actor=admin)
    assert exc.value.code == "IMPORT_JOB_CLOSED"

    again = imp.preview_patients(filename="p.xlsx", content=content, actor=admin)
    assert _rows(again)[2].status == "duplicate"  # registered by the first import
    both = imp.confirm_job(again, actor=admin, include_duplicates=True)
    assert both.imported_rows == 2
    assert Patient.objects.filter(full_name_ar="عثمان موسى").count() == 2


def test_an_age_only_row_is_checked_against_the_estimated_birth_date(admin, clerk) -> None:
    """Review (e2e): the preview skipped name+birth-date matches for age-only rows, so a row
    shown as valid was then refused as a duplicate at confirm."""
    today = date(2026, 10, 8)
    ps.register_patient(
        ps.PatientData(sex="female", full_name_ar="رقية الحاج", age_years=30),
        actor=clerk,
        today=today,
    )
    content = _xlsx([["رقية الحاج", "", "F", None, 30, ""]])
    job = imp.preview_patients(filename="p.xlsx", content=content, actor=admin, today=today)
    row = _rows(job)[2]
    assert row.status == "duplicate"
    assert row.warnings[0]["code"] == "name_dob"


def test_a_file_registered_after_the_preview_is_not_created_twice(admin, clerk) -> None:
    content = _xlsx([["نور الهدى", "", "F", None, 20, "0912777300"]])
    job = imp.preview_patients(filename="p.xlsx", content=content, actor=admin)
    ps.register_patient(
        ps.PatientData(sex="female", full_name_ar="نور الهدى", phone="0912777300"), actor=clerk
    )
    done = imp.confirm_job(job, actor=admin)
    row = _rows(done)[2]
    assert (row.status, row.errors) == ("skipped", [{"code": "DUPLICATE_PATIENT", "field": ""}])
    assert done.imported_rows == 0


def test_csv_with_arabic_headers(admin) -> None:
    text = "الاسم بالعربية,الجنس,العمر (سنوات),الهاتف\nمحمد,ذكر,33,0912777400\n"
    job = imp.preview_patients(filename="p.csv", content=text.encode("utf-8-sig"), actor=admin)
    assert (job.total_rows, job.valid_rows) == (1, 1)
    assert _rows(job)[2].data["age_years"] == 33


@pytest.mark.parametrize(
    ("filename", "content", "code"),
    [
        ("p.pdf", b"%PDF", "IMPORT_FILE_INVALID"),
        ("p.xlsx", b"not a zip", "IMPORT_FILE_INVALID"),
        ("p.csv", b"Phone\n0912\n", "IMPORT_HEADERS_MISSING"),
        ("p.csv", b"Name (English),Sex\n", "IMPORT_NO_ROWS"),
        ("p.csv", b"", "IMPORT_NO_ROWS"),
    ],
)
def test_unreadable_files(admin, filename: str, content: bytes, code: str) -> None:
    with pytest.raises(DomainError) as exc:
        imp.preview_patients(filename=filename, content=content, actor=admin)
    assert exc.value.code == code
    assert not ImportJob.objects.exists()


def test_limits(admin, monkeypatch) -> None:
    monkeypatch.setattr(imp, "MAX_ROWS", 2)
    content = _xlsx([["A", "", "M", None, 1, ""]] * 3)
    with pytest.raises(DomainError) as exc:
        imp.preview_patients(filename="p.xlsx", content=content, actor=admin)
    assert (exc.value.code, exc.value.details["limit"]) == ("IMPORT_TOO_MANY_ROWS", 2)
    monkeypatch.setattr(imp, "MAX_FILE_BYTES", 10)
    with pytest.raises(DomainError) as exc:
        imp.preview_patients(filename="p.xlsx", content=content, actor=admin)
    assert exc.value.code == "IMPORT_FILE_TOO_LARGE"


def test_needs_the_import_permission(clerk, admin) -> None:
    content = _xlsx([["A", "", "M", None, 1, ""]])
    with pytest.raises(PermissionRequired):
        imp.preview_patients(filename="p.xlsx", content=content, actor=clerk)
    job = imp.preview_patients(filename="p.xlsx", content=content, actor=admin)
    with pytest.raises(PermissionRequired):
        imp.confirm_job(job, actor=clerk)
    cancelled = imp.cancel_job(job, actor=admin)
    assert cancelled.status == "cancelled"
    with pytest.raises(DomainError) as exc:
        imp.confirm_job(cancelled, actor=admin)
    assert exc.value.code == "IMPORT_JOB_CLOSED"


def test_template_headers_read_back(admin) -> None:
    for language in ("en", "ar"):
        sheet = load_workbook(io.BytesIO(imp.patient_template(language))).active
        assert sheet is not None
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        job = imp.preview_patients(
            filename="t.xlsx",
            content=_xlsx([["سعاد", "", "F", None, 25, "", "", "", "", "", "", ""]], header),
            actor=admin,
        )
        assert job.valid_rows == 1


# --- API ------------------------------------------------------------------------------------


def _upload(api: ApiClient, name: str, content: bytes) -> Any:
    token = api.csrftoken or api.fetch_csrf()
    return api.django.post(
        "/api/imports/patients",
        {"file": _named(name, content)},
        headers={"X-CSRFToken": token},
    )


def _named(name: str, content: bytes) -> Any:
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(name, content)


def test_api_preview_rows_confirm(make_user, api_client: ApiClient) -> None:
    make_user("boss", roles=["admin"])
    assert api_client.login("boss").status_code == 200
    content = _xlsx([["هبة", "", "F", None, 21, "0912777500"], ["", "", "F", None, 1, ""]])
    created = _upload(api_client, "p.xlsx", content)
    assert created.status_code == 201, created.content
    job = created.json()
    assert (job["valid_rows"], job["error_rows"], job["status"]) == (1, 1, "validated")

    problems = api_client.get(f"/api/imports/{job['id']}/rows?status=problems").json()
    assert [r["row_no"] for r in problems["items"]] == [3]
    assert problems["items"][0]["errors"][0]["code"] == "NAME_REQUIRED"
    every = api_client.get(f"/api/imports/{job['id']}/rows").json()
    assert every["count"] == 2

    done = api_client.post(f"/api/imports/{job['id']}/confirm", {"include_duplicates": False})
    assert done.status_code == 200, done.content
    assert done.json()["imported_rows"] == 1
    again = api_client.post(f"/api/imports/{job['id']}/confirm", {})
    assert again.status_code == 409
    assert again.json()["code"] == "IMPORT_JOB_CLOSED"

    bad = _upload(api_client, "p.txt", b"x")
    assert (bad.status_code, bad.json()["code"]) == (409, "IMPORT_FILE_INVALID")

    template = api_client.get("/api/imports/patients/template?language=ar")
    assert template.status_code == 200
    assert template["Content-Type"].startswith("application/vnd.openxmlformats")


def test_api_refuses_reception(make_user, api_client: ApiClient) -> None:
    make_user("desk", roles=["receptionist"])
    assert api_client.login("desk").status_code == 200
    response = _upload(api_client, "p.xlsx", _xlsx([]))
    assert (response.status_code, response.json()["code"]) == (403, "PERMISSION_DENIED")
    assert api_client.get("/api/imports/patients/template").status_code == 403
