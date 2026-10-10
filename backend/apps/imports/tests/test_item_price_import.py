"""Excel import of stock items with opening stock (FEATURES 8.13) and of price list items into
a version that has not started (FEATURES 5.2): preview, validation, duplicates, confirm.

Opening stock reaches the shelves only through a posted goods receipt (pharmacy services),
never by writing stock rows; prices only through the catalog services, never into a version
that is already effective (invariant 6).
"""

from __future__ import annotations

import io
import zipfile
from datetime import date, timedelta
from decimal import Decimal
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import Workbook

from api.errors import PermissionRequired
from apps.catalog import services as catalog
from apps.catalog.models import PriceItem, PriceList, PriceListVersion, Service
from apps.core.models import Role, RolePermission
from apps.core.tests import builders as b
from apps.imports import services as imp
from apps.imports.models import ImportJob, ImportRow
from apps.pharmacy import services as ps
from apps.pharmacy.models import (
    Batch,
    GoodsReceipt,
    Item,
    StockBalance,
    StockMove,
    UnitConversion,
)
from domain.errors import DomainError

pytestmark = pytest.mark.django_db

ITEM_HEADER = [
    "Service code",
    "Name (English)",
    "Name (Arabic)",
    "Kind (drug or consumable)",
    "Generic name",
    "Form",
    "Base unit code",
    "Pack unit code",
    "Base units per pack",
    "Minimum stock",
    "Batch number",
    "Expiry date",
    "Opening quantity (base units)",
    "Unit cost (per base unit)",
    "Store code",
]


@pytest.fixture(autouse=True)
def _media(settings: Any, tmp_path: Any) -> None:
    settings.MEDIA_ROOT = str(tmp_path)


@pytest.fixture
def admin(make_user):
    return make_user("importer", roles=["admin"])


@pytest.fixture
def pharmacy_store():
    return b.store("PHA")


def xlsx(header: list[str], rows: list[list[Any]]) -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def rows_of(job: ImportJob) -> dict[int, ImportRow]:
    return {r.row_no: r for r in job.rows.all()}


def far(days: int = 400) -> date:
    return timezone.localdate() + timedelta(days=days)


def zinc(code: str = "DRG-ZINC20", **over: Any) -> list[Any]:
    row = {
        "code": code,
        "name_en": "Zinc 20 mg",
        "name_ar": "زنك",
        "kind": "drug",
        "generic": "Zinc sulfate",
        "form": "tablet",
        "base": "TAB",
        "pack": "BOX",
        "factor": 100,
        "min": 50,
        "batch": "Z-001",
        "expiry": far(),
        "qty": 300,
        "cost": "12.5",
        "store": "",
    }
    row.update(over)
    return list(row.values())


# --- items, batches and opening stock -------------------------------------------------------


def test_item_preview_validates_and_flags_existing_stock(admin, pharmacy_store) -> None:
    existing = b.item(generic_name="Paracetamol")
    stocked = b.batch(existing, far(200))
    b.stock_move(stocked, pharmacy_store, "10")
    bare = b.item(generic_name="Ibuprofen")  # an item with no stock yet
    lab = b.service(kind="lab", code="LAB-X")
    content = xlsx(
        ITEM_HEADER,
        [
            zinc(),  # 2: new service, item, unit and batch
            zinc(batch="Z-002", qty=20),  # 3: a second batch of the same item
            zinc(existing.service.code, batch=stocked.batch_no, expiry=stocked.expiry_date),
            zinc(bare.service.code, batch="I-1"),  # 5: existing item, no stock: a warning
            zinc(lab.code),  # 6: not a drug or consumable
            zinc("DRG-NEW2", qty=0),  # 7: bad quantity
            zinc("DRG-NEW3", expiry=timezone.localdate() - timedelta(days=1)),  # 8: expired
        ],
    )
    job = imp.preview(
        "items", filename="items.xlsx", content=content, actor=admin, options={"store": "pha"}
    )
    assert job.kind == "items"
    assert job.options == {"store": "PHA"}
    assert (job.total_rows, job.valid_rows, job.error_rows, job.duplicate_rows) == (7, 3, 3, 1)
    rows = rows_of(job)
    assert [rows[n].status for n in range(2, 9)] == [
        "valid",
        "valid",
        "duplicate",
        "warning",
        "error",
        "error",
        "error",
    ]
    assert rows[2].data["store"] == "PHA"
    assert {w["code"] for w in rows[4].warnings} == {"item_exists", "batch_exists", "stock_exists"}
    assert rows[5].warnings == [{"code": "item_exists"}]
    assert rows[6].errors == [{"code": "SERVICE_NOT_STOCKABLE", "field": "service_code"}]
    assert rows[7].errors == [{"code": "INVALID_QUANTITY", "field": "quantity"}]
    assert rows[8].errors == [{"code": "BATCH_EXPIRED", "field": "expiry_date"}]
    # A preview creates nothing.
    assert not Service.objects.filter(code="DRG-ZINC20").exists()
    assert not GoodsReceipt.objects.exists()


def test_item_confirm_creates_items_and_posts_opening_stock_through_a_receipt(
    admin, pharmacy_store
) -> None:
    main = b.store("MAIN")
    bare = b.item(generic_name="Ibuprofen")
    content = xlsx(
        ITEM_HEADER,
        [
            zinc(),
            zinc(batch="Z-002", qty=20, store="MAIN"),
            zinc(bare.service.code, batch="I-1", qty=40, cost="3"),
            zinc("DRG-BAD", base=""),  # error: never imported
        ],
    )
    job = imp.preview(
        "items", filename="items.xlsx", content=content, actor=admin, options={"store": "PHA"}
    )
    done = imp.confirm_job(job, actor=admin)
    assert done.status == "confirmed"
    assert done.imported_rows == 3
    rows = rows_of(done)
    assert [rows[n].status for n in (2, 3, 4, 5)] == ["imported", "imported", "imported", "error"]

    service = Service.objects.get(code="DRG-ZINC20")
    assert (service.kind, service.name_en) == ("drug", "Zinc 20 mg")
    item = Item.objects.get(service=service)
    assert (item.generic_name, item.base_unit_code, item.min_stock) == (
        "Zinc sulfate",
        "TAB",
        Decimal(50),
    )
    assert UnitConversion.objects.get(item=item).factor == 100
    assert rows[2].result_id == item.pk

    # One posted receipt per store from the opening-stock supplier, never a direct stock row.
    receipts = GoodsReceipt.objects.filter(status="posted").order_by("store__code")
    assert [(r.store.code, r.supplier.code) for r in receipts] == [
        ("MAIN", "OPENING"),
        ("PHA", "OPENING"),
    ]
    assert [r.supplier_invoice_no for r in receipts] == [
        f"IMPORT-{job.pk}-MAIN",
        f"IMPORT-{job.pk}-PHA",
    ]
    moves = StockMove.objects.filter(item__in=[item, bare])
    assert set(moves.values_list("kind", "source_type")) == {("receipt", "receipt_line")}
    assert ps.on_hand(item, pharmacy_store) == 300
    assert ps.on_hand(item, main) == 20
    assert ps.on_hand(bare, pharmacy_store) == 40
    batch = Batch.objects.get(item=bare, batch_no="I-1")
    assert batch.unit_cost == Decimal("3.0000")
    assert batch.supplier is not None
    assert batch.supplier.code == "OPENING"
    assert StockBalance.objects.filter(batch=batch, store=pharmacy_store).get().qty_base == 40
    assert done.summary["receipts"] == sorted(r.number for r in receipts)

    with pytest.raises(DomainError) as closed:
        imp.confirm_job(job, actor=admin)
    assert closed.value.code == "IMPORT_JOB_CLOSED"


def test_possible_duplicates_post_stock_only_when_asked(admin, pharmacy_store) -> None:
    existing = b.item(generic_name="Paracetamol")
    stocked = b.batch(existing, far(200))
    b.stock_move(stocked, pharmacy_store, "10")
    content = xlsx(ITEM_HEADER, [zinc(existing.service.code, batch="P-NEW", qty=5, store="PHA")])
    first = imp.preview("items", filename="a.xlsx", content=content, actor=admin, options={})
    assert rows_of(first)[2].status == "duplicate"  # the item already has stock
    with pytest.raises(DomainError) as nothing:
        imp.confirm_job(first, actor=admin)
    assert nothing.value.code == "IMPORT_NOTHING_TO_IMPORT"
    first.refresh_from_db()
    assert first.status == "validated"
    assert ps.on_hand(existing, pharmacy_store) == 10

    done = imp.confirm_job(first, actor=admin, include_duplicates=True)
    assert done.imported_rows == 1
    assert ps.on_hand(existing, pharmacy_store) == 15


def test_a_store_is_needed_for_opening_stock(admin, pharmacy_store) -> None:
    job = imp.preview(
        "items", filename="a.xlsx", content=xlsx(ITEM_HEADER, [zinc()]), actor=admin, options={}
    )
    assert rows_of(job)[2].errors == [{"code": "STORE_UNKNOWN", "field": "store"}]
    with pytest.raises(DomainError) as bad:
        imp.preview(
            "items",
            filename="a.xlsx",
            content=xlsx(ITEM_HEADER, [zinc()]),
            actor=admin,
            options={"store": "NOPE"},
        )
    assert bad.value.code == "STORE_UNKNOWN"


def test_formula_cells_are_refused_and_never_evaluated(admin, pharmacy_store) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(ITEM_HEADER)
    sheet.append(zinc(qty="=10*30", store="PHA"))
    sheet.append(zinc("DRG-OK", batch="OK-1", store="PHA"))
    sheet["B3"] = '=HYPERLINK("http://evil.example","x")'
    out = io.BytesIO()
    book.save(out)
    job = imp.preview("items", filename="f.xlsx", content=out.getvalue(), actor=admin, options={})
    rows = rows_of(job)
    assert rows[2].errors == [{"code": "FORMULA_NOT_ALLOWED", "field": "quantity"}]
    assert rows[3].errors == [{"code": "FORMULA_NOT_ALLOWED", "field": "name_en"}]
    assert "HYPERLINK" not in str(rows[3].data)


def test_item_import_needs_the_pharmacy_permissions(make_user, pharmacy_store) -> None:
    clerk = make_user("clerk", roles=["receptionist"])
    RolePermission.objects.create(
        role=Role.objects.get(code="receptionist"), code="imports.run", allowed=True
    )
    content = xlsx(ITEM_HEADER, [zinc(store="PHA")])
    with pytest.raises(PermissionRequired):
        imp.preview("items", filename="a.xlsx", content=content, actor=clerk, options={})


def test_new_catalog_codes_need_the_catalog_permission(make_user, pharmacy_store) -> None:
    pharmacist = make_user("pharm-imp", roles=["pharmacist"])
    RolePermission.objects.create(
        role=Role.objects.get(code="pharmacist"), code="imports.run", allowed=True
    )
    existing = b.item(generic_name="Known")
    new_code = imp.preview(
        "items",
        filename="a.xlsx",
        content=xlsx(ITEM_HEADER, [zinc(store="PHA")]),
        actor=pharmacist,
        options={},
    )
    with pytest.raises(PermissionRequired):
        imp.confirm_job(new_code, actor=pharmacist)
    assert not Service.objects.filter(code="DRG-ZINC20").exists()
    known = imp.preview(
        "items",
        filename="b.xlsx",
        content=xlsx(ITEM_HEADER, [zinc(existing.service.code, batch="K-1", store="PHA")]),
        actor=pharmacist,
        options={},
    )
    assert imp.confirm_job(known, actor=pharmacist).imported_rows == 1
    assert ps.on_hand(existing, pharmacy_store) == 300


def test_a_barcode_already_taken_is_an_error_in_the_preview(admin, pharmacy_store) -> None:
    b.item(generic_name="Taken", barcode="111")
    content = xlsx([*ITEM_HEADER, "Barcode"], [[*zinc("DRG-A", batch="A-1"), "111"]])
    job = imp.preview(
        "items", filename="a.xlsx", content=content, actor=admin, options={"store": "PHA"}
    )
    assert rows_of(job)[2].errors == [{"code": "BARCODE_TAKEN", "field": "barcode"}]


def test_an_item_failure_at_confirm_skips_only_its_rows(admin, pharmacy_store) -> None:
    content = xlsx(
        [*ITEM_HEADER, "Barcode"],
        [[*zinc("DRG-A", batch="A-1"), "111"], [*zinc("DRG-B", batch="B-1"), ""]],
    )
    job = imp.preview(
        "items", filename="a.xlsx", content=content, actor=admin, options={"store": "PHA"}
    )
    assert job.valid_rows == 2
    b.item(generic_name="Taken since", barcode="111")  # between preview and confirm
    done = imp.confirm_job(job, actor=admin)
    rows = rows_of(done)
    assert rows[2].status == "skipped"
    assert rows[2].errors == [{"code": "BARCODE_TAKEN", "field": ""}]
    assert rows[3].status == "imported"
    assert not Service.objects.filter(code="DRG-A").exists()  # rolled back with its item
    assert Item.objects.filter(service__code="DRG-B").exists()
    assert ps.on_hand(Item.objects.get(service__code="DRG-B"), pharmacy_store) == 300


# --- reading files safely -------------------------------------------------------------------


def test_files_that_are_not_spreadsheets_are_refused(admin) -> None:
    for name, content in (
        ("a.xlsx", b"not a zip"),
        ("a.exe", b"MZ"),
        ("a.csv", b"Service code,Price\x00\n"),
    ):
        with pytest.raises(DomainError) as bad:
            imp.preview(
                "prices",
                filename=name,
                content=content,
                actor=admin,
                options={"price_list": "CASH", "effective_from": far(5).isoformat()},
            )
        assert bad.value.code == "IMPORT_FILE_INVALID", name


def test_a_decompression_bomb_is_refused_before_parsing(admin, settings) -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"0" * (imp.MAX_UNPACKED_BYTES + 1))
    with pytest.raises(DomainError) as bad:
        imp.preview("items", filename="bomb.xlsx", content=out.getvalue(), actor=admin, options={})
    assert bad.value.code == "IMPORT_FILE_TOO_LARGE"


def test_unknown_kinds_are_refused(admin) -> None:
    with pytest.raises(DomainError) as bad:
        imp.preview("icd10", filename="a.csv", content=b"x", actor=admin, options={})
    assert bad.value.code == "IMPORT_KIND_UNKNOWN"


# --- price list items -----------------------------------------------------------------------

PRICE_HEADER = ["Service code", "Price (SDG)", "Service name (for reference)"]


@pytest.fixture
def priced(admin):
    """Two services on the cash list's effective version (a reference version)."""
    cash = PriceList.objects.get(code="CASH")
    a = b.service(kind="consultation", code="CONS-A")
    c = b.service(kind="lab", code="LAB-C")
    version = b.price_version(cash)
    PriceItem.objects.create(version=version, service=a, unit_price=Decimal("15000.00"))
    PriceItem.objects.create(version=version, service=c, unit_price=Decimal("12000.00"))
    return cash, version, a, c


def price_job(admin, rows: list[list[Any]], effective: date, plist: str = "CASH") -> ImportJob:
    return imp.preview(
        "prices",
        filename="prices.xlsx",
        content=xlsx(PRICE_HEADER, rows),
        actor=admin,
        options={"price_list": plist, "effective_from": effective.isoformat()},
    )


def test_price_preview_shows_the_change_against_the_base_version(admin, priced) -> None:
    job = price_job(
        admin,
        [["CONS-A", 16500, "x"], ["LAB-C", "12000"], ["NOPE", 1], ["CONS-A", "-5"]],
        far(3),
    )
    assert job.options == {"price_list": "CASH", "effective_from": far(3).isoformat()}
    assert (job.valid_rows, job.error_rows) == (2, 2)
    rows = rows_of(job)
    assert rows[2].data["current_price"] == "15000.00"
    assert rows[2].data["price"] == "16500.00"
    assert rows[3].status == "warning"
    assert rows[3].warnings == [{"code": "unchanged"}]
    assert rows[4].errors == [{"code": "SERVICE_UNKNOWN", "field": "service_code"}]


def test_prices_never_go_into_an_effective_version(admin, priced) -> None:
    for day in (timezone.localdate(), timezone.localdate() - timedelta(days=1)):
        with pytest.raises(DomainError) as bad:
            price_job(admin, [["CONS-A", 1]], day)
        assert bad.value.code == "PRICE_VERSION_BACKDATED"
    for options, code in (
        ({}, "IMPORT_OPTION_REQUIRED"),
        ({"price_list": "NOPE", "effective_from": far(2).isoformat()}, "PRICE_LIST_UNKNOWN"),
        ({"price_list": "CASH", "effective_from": "soon"}, "IMPORT_OPTION_INVALID"),
    ):
        with pytest.raises(DomainError) as refused:
            imp.preview(
                "prices",
                filename="p.xlsx",
                content=xlsx(PRICE_HEADER, [["CONS-A", 1]]),
                actor=admin,
                options=options,
            )
        assert refused.value.code == code


def test_price_confirm_creates_a_future_version_and_leaves_the_effective_one(admin, priced) -> None:
    cash, version, a, c = priced
    job = price_job(admin, [["CONS-A", 16500], ["NOPE", 1]], far(3))
    done = imp.confirm_job(job, actor=admin)
    assert done.imported_rows == 1
    new = PriceListVersion.objects.get(price_list=cash, effective_from=far(3))
    assert catalog.version_prices(new) == {a.pk: Decimal("16500.00"), c.pk: Decimal("12000.00")}
    assert catalog.version_prices(version) == {a.pk: Decimal("15000.00"), c.pk: Decimal("12000.00")}
    assert rows_of(done)[2].result_id == new.pk
    assert done.summary["version_id"] == new.pk


def test_price_confirm_edits_a_version_that_has_not_started(admin, priced) -> None:
    cash, _, a, c = priced
    planned = catalog.derive_version(cash, effective_from=far(4), actor=admin, changes={})
    done = imp.confirm_job(price_job(admin, [["LAB-C", "13000.5"]], far(4)), actor=admin)
    assert done.status == "confirmed"
    assert PriceListVersion.objects.filter(price_list=cash, effective_from=far(4)).count() == 1
    assert catalog.version_prices(planned)[c.pk] == Decimal("13000.50")
    assert catalog.version_prices(planned)[a.pk] == Decimal("15000.00")


def test_price_import_needs_the_price_permission(make_user, priced) -> None:
    clerk = make_user("clerk2", roles=["receptionist"])
    RolePermission.objects.create(
        role=Role.objects.get(code="receptionist"), code="imports.run", allowed=True
    )
    with pytest.raises(PermissionRequired):
        price_job(clerk, [["CONS-A", 1]], far(3))


# --- templates and the patient import through the same flow ---------------------------------


@pytest.mark.parametrize("kind", ["patients", "items", "prices"])
@pytest.mark.parametrize("language", ["ar", "en"])
def test_templates_have_the_header_row(kind: str, language: str) -> None:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(imp.template(kind, language)), read_only=True)
    header = next(book.worksheets[0].iter_rows(values_only=True))
    assert header
    assert all(isinstance(h, str) and h for h in header)


def test_patients_preview_through_the_generic_entry(admin) -> None:
    content = xlsx(
        ["Name (Arabic)", "Sex", "Phone"], [["سارة", "F", "0912000555"], ["=1+1", "M", ""]]
    )
    job = imp.preview("patients", filename="p.xlsx", content=content, actor=admin, options={})
    rows = rows_of(job)
    assert job.kind == "patients"
    assert rows[2].status == "valid"
    assert {"code": "FORMULA_NOT_ALLOWED", "field": "full_name_ar"} in rows[3].errors


# --- API ------------------------------------------------------------------------------------


def _post_job(api: Any, name: str, content: bytes, **fields: str) -> Any:
    from django.core.files.uploadedfile import SimpleUploadedFile

    token = api.csrftoken or api.fetch_csrf()
    return api.django.post(
        "/api/imports/jobs",
        {"file": SimpleUploadedFile(name, content), **fields},
        headers={"X-CSRFToken": token},
    )


def test_api_items_preview_rows_confirm_and_history(make_user, api_client, pharmacy_store) -> None:
    make_user("boss", roles=["admin"])
    assert api_client.login("boss").status_code == 200
    content = xlsx(ITEM_HEADER, [zinc(), zinc("DRG-X", base="")])
    created = _post_job(api_client, "items.xlsx", content, kind="items", store="PHA")
    assert created.status_code == 201, created.content
    job = created.json()
    assert (job["kind"], job["valid_rows"], job["error_rows"]) == ("items", 1, 1)
    assert job["options"] == {"store": "PHA"}

    rows = api_client.get(f"/api/imports/jobs/{job['id']}/rows?status=problems").json()
    assert [r["row_no"] for r in rows["items"]] == [3]
    assert rows["items"][0]["data"]["service_code"] == "DRG-X"
    assert rows["items"][0]["errors"] == [{"code": "UNIT_REQUIRED", "field": "base_unit_code"}]

    done = api_client.post(f"/api/imports/{job['id']}/confirm", {"include_duplicates": False})
    assert done.status_code == 200, done.content
    body = done.json()
    assert body["imported_rows"] == 1
    assert len(body["summary"]["receipts"]) == 1

    history = api_client.get("/api/imports/jobs?kind=items").json()
    assert [j["id"] for j in history["items"]] == [job["id"]]
    template = api_client.get("/api/imports/templates/prices?language=ar")
    assert template.status_code == 200
    assert template["Content-Disposition"] == 'attachment; filename="prices-ar.xlsx"'
    assert template["X-Content-Type-Options"] == "nosniff"


def test_api_price_options_and_errors(make_user, api_client, priced) -> None:
    make_user("boss2", roles=["admin"])
    assert api_client.login("boss2").status_code == 200
    content = xlsx(PRICE_HEADER, [["CONS-A", 20000]])
    missing = _post_job(api_client, "p.xlsx", content, kind="prices")
    assert (missing.status_code, missing.json()["code"]) == (409, "IMPORT_OPTION_REQUIRED")
    unknown_kind = _post_job(api_client, "p.xlsx", content, kind="icd10")
    assert unknown_kind.status_code == 422
    ok = _post_job(
        api_client,
        "p.xlsx",
        content,
        kind="prices",
        price_list="CASH",
        effective_from=far(6).isoformat(),
    )
    assert ok.status_code == 201, ok.content
    assert ok.json()["options"]["effective_from"] == far(6).isoformat()


def test_api_refuses_without_the_permission(make_user, api_client) -> None:
    make_user("desk2", roles=["receptionist"])
    assert api_client.login("desk2").status_code == 200
    response = _post_job(api_client, "p.xlsx", b"x", kind="items")
    assert (response.status_code, response.json()["code"]) == (403, "PERMISSION_DENIED")
    assert api_client.get("/api/imports/jobs").status_code == 403
    assert api_client.get("/api/imports/templates/items").status_code == 403
