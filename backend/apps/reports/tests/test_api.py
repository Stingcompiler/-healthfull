"""``/api/reports``: every report as JSON and Excel, permissions, filters and labels."""

from __future__ import annotations

import io
from datetime import timedelta
from decimal import Decimal
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import User
from apps.core.permissions import PERMISSIONS, effective_permissions
from apps.payments.tests import api_kit as kit
from apps.reports import labels
from apps.reports.registry import REPORTS
from apps.reports.tests.scenario import Day
from conftest import router_operations

pytestmark = pytest.mark.django_db

FINANCE = [
    k
    for k, s in REPORTS.items()
    if s.permission in ("reports.view_finance", "reports.view_exceptions")
]


def _login(user: User) -> Any:
    user.set_password(kit.TEST_PASSWORD)
    user.save()
    client = kit.ApiClient()
    assert client.login(user.username, kit.TEST_PASSWORD).status_code == 200
    return client


@pytest.fixture
def manager_api(day: Day) -> Any:
    return _login(day.manager)


def test_the_router_exposes_one_read_and_one_export_per_report() -> None:
    ops = router_operations("/reports")
    expected: set[tuple[str, str, str, str | None]] = {("GET", "/ping", "reports_get_ping", None)}
    expected.add(("GET", "/dashboard", "reports_get_dashboard", "reports.view_dashboard"))
    for key, spec in REPORTS.items():
        expected.add(("GET", f"/{key}", f"reports_get_{key}", spec.permission))
        expected.add(("GET", f"/{key}/export", f"reports_export_{key}", spec.permission))
    assert ops == expected


def test_permission_defaults() -> None:
    def roles(code: str) -> set[str]:
        return set(PERMISSIONS[code].default_roles)

    finance = {"cashier_supervisor", "accountant", "manager", "admin"}
    assert roles("reports.view_finance") == finance
    assert roles("reports.view_exceptions") == finance
    assert roles("reports.view_dashboard") == {"manager", "admin"}
    assert roles("reports.view_stock") == {"pharmacist", "accountant", "manager", "admin"}
    assert roles("reports.view_visits") == {"manager", "admin"}
    assert roles("reports.view_lab") == {"lab_supervisor", "manager", "admin"}
    for code in [c for c in PERMISSIONS if c.startswith("reports.")]:
        assert not ({"doctor", "cashier", "receptionist", "nurse"} & roles(code)), code


def test_every_report_answers_with_labels_in_both_languages(day: Day, manager_api: Any) -> None:
    for key in REPORTS:
        body = kit.ok(manager_api.get(f"/api/reports/{key}"))
        assert body["key"] == key
        assert all(body["title"].values())
        for section in body["sections"]:
            assert all(section["label"].values()), (key, section["key"])
            for column in section["columns"]:
                assert all(column["label"].values()), (key, column["key"])
                for row in section["rows"]:
                    assert set(row) <= {c["key"] for c in section["columns"]}
        for metric in body["metrics"]:
            assert all(metric["label"].values()), (key, metric["key"])


def test_revenue_json_carries_money_as_strings(day: Day, manager_api: Any) -> None:
    today = timezone.localdate().isoformat()
    body = kit.ok(manager_api.get(f"/api/reports/revenue?date_from={today}&date_to={today}"))
    metrics = {m["key"]: m["value"] for m in body["metrics"]}
    assert metrics["net_revenue"] == "21500.00"
    assert metrics["pending_transfers"] == "6000.00"
    assert metrics["collected"] == "13500.00"
    assert body["filters"] == {
        "date_from": today,
        "date_to": today,
        "department_id": None,
        "user_id": None,
        "days": None,
    }
    assert body["filters_available"] == ["dates", "department"]
    assert {d["id"] for d in body["departments"]} >= {day.opd.pk, day.lab.pk}


def test_filters_a_report_does_not_take_are_dropped(day: Day, manager_api: Any) -> None:
    body = kit.ok(manager_api.get(f"/api/reports/revenue?user_id={day.cashier.pk}"))
    assert body["filters"]["user_id"] is None
    body = kit.ok(manager_api.get("/api/reports/stock_expiry"))
    assert body["filters"]["days"] == 90


def test_date_errors(day: Day, manager_api: Any) -> None:
    today = timezone.localdate()
    kit.error(
        manager_api.get(
            f"/api/reports/revenue?date_from={today}&date_to={today - timedelta(days=1)}"
        ),
        409,
        "INVALID_DATE_RANGE",
    )
    body = kit.error(
        manager_api.get(
            f"/api/reports/revenue?date_from={today - timedelta(days=400)}&date_to={today}"
        ),
        409,
        "REPORT_RANGE_TOO_LONG",
    )
    assert body["details"]["max_days"] == 366


def test_excel_export_is_a_valid_workbook(day: Day, manager_api: Any) -> None:
    today = timezone.localdate().isoformat()
    response = manager_api.get(
        f"/api/reports/revenue/export?date_from={today}&date_to={today}&language=en"
    )
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert f'filename="revenue-{today}-en.xlsx"' in response["Content-Disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "Summary"
    summary = {
        row[0]: row[1] for row in book["Summary"].iter_rows(values_only=True) if row and row[0]
    }
    assert Decimal(str(summary["Net revenue"])) == Decimal("21500.00")
    dept = book["By department"]
    header = [c.value for c in dept[3]]
    assert header[:2] == ["Department", "Gross"]
    total_row = [r for r in dept.iter_rows(values_only=True) if r and r[0] == "Total"]
    assert Decimal(str(total_row[0][4])) == Decimal("21500.00")
    # Money cells carry the two-decimal number format.
    assert dept.cell(row=4, column=2).number_format == "#,##0.00"


def test_arabic_export_reads_right_to_left(day: Day, manager_api: Any) -> None:
    response = manager_api.get("/api/reports/shift_variances/export?language=ar")
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames[0] == "الملخص"
    assert all(ws.sheet_view.rightToLeft for ws in book.worksheets)


def test_every_report_exports(day: Day, manager_api: Any) -> None:
    for key in REPORTS:
        response = manager_api.get(f"/api/reports/{key}/export?language=en")
        assert response.status_code == 200, (key, response.content[:200])
        load_workbook(io.BytesIO(response.content))


@pytest.mark.parametrize("typed", ['=HYPERLINK("http://x","y")', "+1+2", "-3+4", "@SUM(A1)"])
def test_export_never_writes_formulas(day: Day, manager_api: Any, typed: str) -> None:
    """A name typed as a formula comes out as quoted text (spreadsheet formula injection)."""
    from apps.patients.models import Patient

    Patient.objects.filter(pk=day.patient_a.pk).update(full_name_en=typed)
    response = manager_api.get("/api/reports/pending_transfers/export?language=en")
    book = load_workbook(io.BytesIO(response.content))
    cells = [c for ws in book.worksheets for row in ws.iter_rows() for c in row]
    assert not [c for c in cells if c.data_type == "f"]
    (name,) = [c for c in cells if c.value == typed]
    assert name.data_type == "s"
    assert name.quotePrefix


def test_dashboard(day: Day, manager_api: Any) -> None:
    body = kit.ok(manager_api.get("/api/reports/dashboard"))
    metrics = {m["key"]: m["value"] for m in body["metrics"]}
    assert metrics["collected_today"] == "13500.00"
    assert metrics["pending_transfers"] == "6000.00"
    assert len(body["trend"]) == 14
    assert body["trend"][-1]["collected"] == "13500.00"
    assert {a["key"] for a in body["alerts"]} >= {"unreviewed_variances"}


def test_doctors_and_cashiers_are_refused_every_finance_report(day: Day) -> None:
    doctor = _login(day.doctor.user)
    cashier = _login(day.cashier)
    for who, client in (("doctor", doctor), ("cashier", cashier)):
        for key in [*FINANCE, "dashboard"]:
            body = kit.error(client.get(f"/api/reports/{key}"), 403, "PERMISSION_DENIED")
            assert body["details"]["permission"].startswith("reports."), (who, key)
            if key != "dashboard":
                kit.error(client.get(f"/api/reports/{key}/export"), 403, "PERMISSION_DENIED")


def test_every_operation_refuses_roles_without_its_code(make_user: Any) -> None:
    """A sweep over the whole router: a role without the operation's code gets 403 before
    anything runs (no data needed)."""
    roles = (
        "receptionist",
        "doctor",
        "cashier",
        "cashier_supervisor",
        "pharmacist",
        "lab_tech",
        "lab_supervisor",
        "nurse",
        "accountant",
        "manager",
    )
    ops = [op for op in router_operations("/reports") if op[3] is not None]
    for role in roles:
        user = make_user(f"probe_{role}", roles=[role])
        held = effective_permissions(user)
        client = _login(user)
        for _, path, op_id, code in ops:
            if code in held:
                continue
            response = client.get(f"/api/reports{path}")
            assert response.status_code == 403, (role, op_id)
            assert response.json()["details"]["permission"] == code


def test_unauthenticated_is_401() -> None:
    client = kit.ApiClient()
    assert client.get("/api/reports/revenue").status_code == 401


def test_every_label_key_exists_in_both_languages() -> None:
    for key, (ar, en) in labels.LABELS.items():
        assert ar.strip(), key
        assert en.strip(), key
    for key in REPORTS:
        assert f"report.{key}" in labels.LABELS
