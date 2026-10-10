"""Any report as an Excel workbook (FEATURES 12.11).

A summary sheet (center, report, filters, metrics), then one sheet per section with its
header, rows and totals. Money is written as numbers with two decimals from the report's
``Decimal`` values (never a float computed here), dates as dates, names in the chosen
language. Arabic workbooks read right to left.

Spreadsheet formula injection: names, notes and references are typed by people, so any text
starting with ``=``, ``+``, ``-``, ``@``, a tab or a carriage return is stored as plain text
with a quote prefix; this export never writes a formula.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from django.utils import timezone

from apps.reports import labels
from apps.reports.kit import MAX_ROWS, Cell, Column, Report, Section, center_name
from apps.reports.labels import Lang

__all__ = ["workbook"]

_MONEY_FORMAT = "#,##0.00"
_PERCENT_FORMAT = "0.0"
_DATE_FORMAT = "yyyy-mm-dd"
_DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
_RISKY = ("=", "+", "-", "@", "\t", "\r")
_SHEET_BAD = set("[]:*?/\\")


def _pick(value: dict[str, str], lang: Lang) -> str:
    return (
        (value.get("ar") or value.get("en") or "")
        if lang == "ar"
        else (value.get("en") or value.get("ar") or "")
    )


def _put(sheet: Any, row: int, col: int, value: Cell, kind: str, lang: Lang) -> Any:
    cell = sheet.cell(row=row, column=col)
    if value is None:
        return cell
    if isinstance(value, dict):
        value = _pick(value, lang)
    if isinstance(value, datetime):
        cell.value = timezone.localtime(value).replace(tzinfo=None)
        cell.number_format = _DATETIME_FORMAT
    elif isinstance(value, date):
        cell.value = value
        cell.number_format = _DATE_FORMAT
    elif isinstance(value, Decimal):
        cell.value = value
        cell.number_format = _PERCENT_FORMAT if kind == "percent" else _MONEY_FORMAT
    elif isinstance(value, int):
        cell.value = value
    else:
        text = str(value)
        cell.value = text
        if text.startswith(_RISKY):
            cell.data_type = "s"
            cell.quotePrefix = True
    return cell


def _sheet_title(text: str, used: set[str]) -> str:
    clean = "".join("-" if ch in _SHEET_BAD else ch for ch in text).strip("'") or "Sheet"
    base = clean[:28]
    title, n = base, 2
    while title.lower() in used:
        title = f"{base[:25]} ({n})"
        n += 1
    used.add(title.lower())
    return title


def _filters(report: Report, lang: Lang, available: tuple[str, ...]) -> list[tuple[str, str]]:
    f = report.filters
    rows: list[tuple[str, str]] = []
    if "dates" in available:
        to = labels.text("export.to", lang)
        rows.append(
            (
                labels.text("export.period", lang),
                f"{f.date_from.isoformat()} {to} {f.date_to.isoformat()}",
            )
        )
    if "as_of" in available:
        rows.append((labels.text("export.as_of", lang), f.date_to.isoformat()))
    all_ = labels.text("export.all", lang)
    if "department" in available:
        dept = next((o.name for o in report.departments if o.id == f.department_id), None)
        rows.append((labels.text("export.department", lang), _pick(dept, lang) if dept else all_))
    if "user" in available:
        user = next((o.name for o in report.users if o.id == f.user_id), None)
        rows.append((labels.text("export.user", lang), _pick(user, lang) if user else all_))
    if "days" in available and f.days is not None:
        rows.append((labels.text("export.days", lang), str(f.days)))
    return rows


def _write_section(sheet: Any, section: Section, lang: Lang) -> None:
    from openpyxl.styles import Alignment, Font

    bold = Font(bold=True)
    sheet.cell(row=1, column=1, value=labels.text(f"section.{section.key}", lang)).font = Font(
        bold=True, size=12
    )
    head = 3
    columns: list[Column] = section.columns
    for col, column in enumerate(columns, start=1):
        cell = sheet.cell(row=head, column=col, value=labels.text(f"column.{column.key}", lang))
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        width = 30 if column.kind in ("name", "text") else 16
        sheet.column_dimensions[cell.column_letter].width = width
    row = head
    for data in section.rows:
        row += 1
        for col, column in enumerate(columns, start=1):
            _put(sheet, row, col, data.get(column.key), column.kind, lang)
    if section.totals is not None:
        row += 1
        sheet.cell(row=row, column=1, value=labels.text("export.totals", lang)).font = bold
        for col, column in enumerate(columns, start=1):
            if col == 1 or column.key not in section.totals:
                continue
            _put(sheet, row, col, section.totals[column.key], column.kind, lang).font = bold
    if section.truncated:
        note = labels.text("export.truncated", lang).format(rows=MAX_ROWS)
        sheet.cell(row=row + 2, column=1, value=note)
    sheet.freeze_panes = sheet.cell(row=head + 1, column=1)


def workbook(report: Report, available: tuple[str, ...], language: str = "ar") -> tuple[str, bytes]:
    """``report`` as an ``.xlsx`` in ``language`` (ar or en), with its download file name."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    lang: Lang = "ar" if language == "ar" else "en"
    book = Workbook()
    summary = book.active or book.create_sheet()
    used: set[str] = set()
    summary.title = _sheet_title(labels.text("export.summary", lang), used)
    summary.sheet_view.rightToLeft = lang == "ar"
    bold = Font(bold=True)
    _put(summary, 1, 1, center_name(), "text", lang)
    summary.cell(row=1, column=1).font = Font(bold=True, size=14)
    summary.cell(row=2, column=1, value=labels.text(f"report.{report.key}", lang)).font = Font(
        bold=True, size=12
    )
    row = 3
    for label, value in _filters(report, lang, available):
        row += 1
        summary.cell(row=row, column=1, value=label).font = bold
        _put(summary, row, 2, value, "text", lang)
    row += 1
    summary.cell(row=row, column=1, value=labels.text("export.generated_at", lang)).font = bold
    _put(summary, row, 2, report.generated_at, "datetime", lang)
    row += 1
    for metric in report.metrics:
        row += 1
        summary.cell(row=row, column=1, value=labels.text(f"metric.{metric.key}", lang))
        _put(summary, row, 2, metric.value, metric.kind, lang)
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 24

    for section in report.sections:
        sheet = book.create_sheet(_sheet_title(labels.text(f"section.{section.key}", lang), used))
        sheet.sheet_view.rightToLeft = lang == "ar"
        _write_section(sheet, section, lang)

    # Belt and braces: nothing in the book may be a formula (see the module docstring).
    for sheet in book.worksheets:
        for cells in sheet.iter_rows():
            for cell in cells:
                if cell.data_type == "f":
                    cell.data_type = "s"
                    cell.quotePrefix = True

    out = io.BytesIO()
    book.save(out)
    f = report.filters
    stamp = f.date_to.isoformat() if f.date_from == f.date_to else f"{f.date_from}_{f.date_to}"
    return f"{report.key}-{stamp}-{lang}.xlsx", out.getvalue()
